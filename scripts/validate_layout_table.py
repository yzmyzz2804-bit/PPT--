"""校验人工审核的版式表与配置映射；不自动接受变更或重建基线。"""
from pathlib import Path
import json
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]

def lookup(data,path):
    for part in path.split('.'):
        if not isinstance(data,dict) or part not in data:
            raise ValueError('配置字段不存在: '+path)
        data=data[part]
    return data

def validate_layout_table(root=ROOT):
    root=Path(root)
    m=json.loads((root/'config/layout_table_mapping.json').read_text(encoding='utf-8'))
    config=json.loads((root/m['source_config']).read_text(encoding='utf-8'))
    workbook=load_workbook(root/m['workbook'],read_only=True,data_only=False)
    try:
        ws=workbook[m['worksheet']]
        if (ws.max_row,ws.max_column)!=(m['expected_rows'],m['expected_columns']):
            raise ValueError('表格行列数不匹配')
        headers=[c.value for c in next(ws.iter_rows(min_row=m['header_row'],max_row=m['header_row']))]
        if headers!=m['headers']: raise ValueError('表头不匹配')
        fields=m['fields']
        if [x['row'] for x in fields]!=list(range(m['header_row']+1,m['expected_rows']+1)):
            raise ValueError('映射必须完整覆盖每一行，不能遗漏或重复')
        names=[x['field'] for x in fields]
        if len(set(names))!=len(names): raise ValueError('规范字段重复')
        bindings=0; suggestions=0
        allrows=list(ws.iter_rows(min_row=m['header_row']+1,values_only=True))
        for field,row in zip(fields,allrows):
            if row[1]!=field['field']: raise ValueError('字段名称或顺序变化: '+field['field'])
            if set(field['cells'])!=set('ABCDEFGH'): raise ValueError('文字快照列不完整')
            for col,expected in field['cells'].items():
                actual=row[ord(col)-65]
                if actual!=expected: raise ValueError(f'表格文字漂移 {col}{field["row"]}: {field["field"]}')
            if not field['config_bindings']: raise ValueError('规范缺配置关联')
            for binding in field['config_bindings']:
                if lookup(config,binding['path'])!=binding['expected']:
                    raise ValueError('配置与已审核表格基线不一致: '+binding['path'])
                bindings+=1
            for path in field['suggestion_only_absent_paths']:
                try: lookup(config,path)
                except ValueError: pass
                else: raise ValueError('建议字段已启用但对照表未更新: '+path)
                suggestions+=1
    finally:
        workbook.close()
    return {'status':'pass','specification_rows':len(fields),'configuration_bindings':bindings,'inactive_suggestion_paths':suggestions,'workbook':m['workbook'],'scope':'文字、结构及配置快照一致性；非PPT视觉或事实验收'}

if __name__=='__main__':
    try: print(json.dumps(validate_layout_table(),ensure_ascii=False,indent=2))
    except (ValueError,KeyError,OSError) as e:
        print(json.dumps({'status':'fail','error':str(e)},ensure_ascii=False)); raise SystemExit(2)
