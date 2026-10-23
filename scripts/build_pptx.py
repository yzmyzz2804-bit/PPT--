"""
build_pptx.py - HTML → PPTX (占位骨架,后续完善)

依赖:
    pip install python-pptx

策略说明:
    HTML→PPTX 是有损转换,无法 100% 还原 CSS。
    本骨架阶段采用"截图法":对每页 .slide 用 Playwright 截图,
    再将图片作为整页背景插入 PPTX (16:9 尺寸)。
    优点:视觉 100% 保真;缺点:PPTX 内文字不可编辑。

    如需可编辑 PPTX,后续可改用解析 HTML 结构的方式
    (按版式映射到 python-pptx 的 placeholder 类型),
    但版式受限,需要逐模板适配。

使用:
    python build_pptx.py                  # 默认输出 build/deck.pptx
    python build_pptx.py --output out.pptx
"""
import sys
import argparse
import yaml
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
DECK_PATH = PROJECT_ROOT / 'deck.yaml'
SLIDES_DIR = PROJECT_ROOT / 'slides'
BUILD_DIR = PROJECT_ROOT / 'build'


def screenshot_slides(deck: dict) -> list:
    """对每页 .slide 截图,返回图片路径列表"""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print('✗ 需要 playwright: pip install playwright && playwright install chromium')
        sys.exit(1)

    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    shots_dir = BUILD_DIR / 'screenshots'
    shots_dir.mkdir(exist_ok=True)

    slides = deck.get('slides', [])
    paths = []
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=True, channel="chrome")
        except Exception:
            browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(viewport={'width': 1920, 'height': 1080})
        for s in slides:
            page_path = SLIDES_DIR / f'{s["id"]}.html'
            if not page_path.exists():
                print(f'  ⚠ 跳过未渲染页: {s["id"]}')
                continue
            page = ctx.new_page()
            page.goto(page_path.as_uri())
            img_path = shots_dir / f'{s["id"]}.png'
            page.screenshot(path=str(img_path), full_page=False)
            paths.append(img_path)
            page.close()
            print(f'  ✓ 截图: {img_path}')
        browser.close()
    return paths


def build_pptx(image_paths: list, output_path: Path, deck: dict):
    """把每张截图整页插入 PPTX (16:9)"""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Emu
    except ImportError:
        print('✗ 需要 python-pptx: pip install python-pptx')
        sys.exit(1)

    prs = Presentation()
    # 16:9 = 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # blank

    for img in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(img),
            left=Emu(0), top=Emu(0),
            width=prs.slide_width, height=prs.slide_height,
        )

    prs.save(str(output_path))
    print(f'✓ PPTX 已生成: {output_path}')


def main():
    ap = argparse.ArgumentParser(description='构建 PPTX (截图法,占位实现)')
    ap.add_argument('--output', default='', help='输出 PPTX 路径')
    args = ap.parse_args()

    if not DECK_PATH.exists():
        print(f'✗ deck.yaml 不存在')
        sys.exit(1)

    with open(DECK_PATH, 'r', encoding='utf-8') as f:
        deck = yaml.safe_load(f)

    print('Step 1: 截图每页...')
    imgs = screenshot_slides(deck)
    if not imgs:
        print('✗ 没有可截图的页面,请先 render_slide.py --all')
        sys.exit(1)

    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    out = Path(args.output) if args.output else BUILD_DIR / f'{deck.get("project", {}).get("name", "deck")}.pptx'

    print('Step 2: 组装 PPTX...')
    build_pptx(imgs, out, deck)

    from datetime import datetime
    deck.setdefault('build', {})
    deck['build']['last_build_pptx'] = str(out)
    with open(DECK_PATH, 'w', encoding='utf-8') as f:
        yaml.dump(deck, f, allow_unicode=True, sort_keys=False, default_flow_style=False)


if __name__ == '__main__':
    main()
